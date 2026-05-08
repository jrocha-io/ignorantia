"""Tests v2.23.1 (Fix 14 of RS-42 dogfood remediation) — Markdown in main loop.

The dogfood RS-42 manuscripts had Markdown markers (**bold**, *italic*,
``[text](url)``) appearing literally in the rendered PDFs and DOCX. The
dogfood diagnosis identified the root cause: ``render_latex.py:335`` and
``render_docx_abnt.py:265`` were calling ``_escape_latex(text)`` /
``doc.add_paragraph(text)`` directly in the main section loop, bypassing
the ``_markdown_to_latex`` helper that already existed for the contextual
preamble path.

Fix 14 routes body and long_quote paragraphs through Markdown processing:

* ``render_latex.py``: body/long_quote paragraphs go through
  ``_markdown_to_latex()`` (which produces ``\\textbf{}`` / ``\\textit{}`` /
  ``\\href{}{}``). References continue to use ``_escape_latex`` because they
  are already-formatted ABNT strings.
* ``render_docx_abnt.py``: new helper ``_add_runs_with_markdown()`` parses
  Markdown markers and adds python-docx ``add_run()`` calls with
  ``bold=True`` / ``italic=True`` flags. References stay plain.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))

import render_latex  # noqa: E402  — sys.path injection above


# ---------------------------------------------------------------------------
# render_latex: body / long_quote / reference behaviour
# ---------------------------------------------------------------------------


def _content(text: str, ptype: str = "body") -> dict:
    return {
        "sections": [
            {
                "id": "01",
                "title": "Discussão",
                "paragraphs": [{"text": text, "type": ptype}],
            },
        ],
        "references": [],
    }


def test_body_paragraph_bold_becomes_textbf(tmp_path: Path) -> None:
    out = tmp_path / "manuscript.tex"
    render_latex.render_tex(
        _content("This claim is **strongly supported** by the data."),
        out,
        title="T",
    )
    text = out.read_text(encoding="utf-8")
    assert "\\textbf{strongly supported}" in text
    # No literal asterisks survive in the body.
    assert "**strongly supported**" not in text


def test_body_paragraph_italic_becomes_textit(tmp_path: Path) -> None:
    out = tmp_path / "manuscript.tex"
    render_latex.render_tex(
        _content("The term *gamification* is widely used."),
        out,
        title="T",
    )
    text = out.read_text(encoding="utf-8")
    assert "\\textit{gamification}" in text
    assert "*gamification*" not in text


def test_body_paragraph_link_becomes_href(tmp_path: Path) -> None:
    out = tmp_path / "manuscript.tex"
    render_latex.render_tex(
        _content("See [the protocol](https://example.org/protocol) for details."),
        out,
        title="T",
    )
    text = out.read_text(encoding="utf-8")
    assert "\\href{https://example.org/protocol}{the protocol}" in text


def test_long_quote_also_processes_markdown(tmp_path: Path) -> None:
    out = tmp_path / "manuscript.tex"
    render_latex.render_tex(
        _content("This is a *quoted* passage with **emphasis**.", ptype="long_quote"),
        out,
        title="T",
    )
    text = out.read_text(encoding="utf-8")
    assert "\\begin{longquote}" in text
    assert "\\textit{quoted}" in text
    assert "\\textbf{emphasis}" in text
    assert "*quoted*" not in text
    assert "**emphasis**" not in text


def test_reference_paragraph_does_not_process_markdown(tmp_path: Path) -> None:
    """Reference type stays plain — ABNT references shouldn't have Markdown."""
    out = tmp_path / "manuscript.tex"
    render_latex.render_tex(
        _content("LITERAL *star* and **double-star** in a reference.", ptype="reference"),
        out,
        title="T",
    )
    text = out.read_text(encoding="utf-8")
    # Reference type is escaped only — markdown markers stay literal because
    # the input is presumed to be an already-formatted ABNT string.
    assert "\\textbf{double-star}" not in text
    assert "\\textit{star}" not in text


def test_special_chars_still_escaped_with_markdown(tmp_path: Path) -> None:
    """Markdown-converted text still escapes LaTeX special chars in the surrounding plain text."""
    out = tmp_path / "manuscript.tex"
    render_latex.render_tex(
        _content("Cost: 50% & coverage **complete**."),
        out,
        title="T",
    )
    text = out.read_text(encoding="utf-8")
    assert "50\\%" in text
    assert "\\&" in text
    assert "\\textbf{complete}" in text


def test_rs42_dogfood_phrase_no_longer_renders_literal_asterisks(tmp_path: Path) -> None:
    """Phrases lifted from RS-42 v1.0.0 with Markdown leakage now render correctly."""
    rs42_phrase = (
        "A maioria dos estudos (16/24) reporta **pontuação progressiva** como "
        "elemento central, com *feedback imediato* convergindo entre todos."
    )
    out = tmp_path / "manuscript.tex"
    render_latex.render_tex(_content(rs42_phrase), out, title="RS-42 v1.0.1")
    text = out.read_text(encoding="utf-8")
    assert "\\textbf{pontuação progressiva}" in text
    assert "\\textit{feedback imediato}" in text
    assert "**pontuação progressiva**" not in text
    assert "*feedback imediato*" not in text


# ---------------------------------------------------------------------------
# render_docx_abnt: same behaviour via _add_runs_with_markdown helper
# ---------------------------------------------------------------------------

try:
    import docx as _docx_check  # noqa: F401

    HAS_DOCX = True
except ImportError:  # pragma: no cover
    HAS_DOCX = False


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_docx_helper_parses_bold_into_run() -> None:
    import docx as docx_mod

    import render_docx_abnt

    doc = docx_mod.Document()
    p = doc.add_paragraph()
    render_docx_abnt._add_runs_with_markdown(p, "Cost is **high** here.")
    runs = p.runs
    # Expect 3 runs: "Cost is ", "high" (bold), " here."
    assert len(runs) == 3
    assert runs[0].text == "Cost is "
    assert runs[1].text == "high"
    assert runs[1].bold is True
    assert runs[2].text == " here."


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_docx_helper_parses_italic_into_run() -> None:
    import docx as docx_mod

    import render_docx_abnt

    doc = docx_mod.Document()
    p = doc.add_paragraph()
    render_docx_abnt._add_runs_with_markdown(p, "The term *gamification* applies.")
    runs = p.runs
    assert len(runs) == 3
    assert runs[1].text == "gamification"
    assert runs[1].italic is True


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_docx_helper_parses_link_text_only() -> None:
    """Links render as text (URL stripped); python-docx hyperlinks require xml plumbing."""
    import docx as docx_mod

    import render_docx_abnt

    doc = docx_mod.Document()
    p = doc.add_paragraph()
    render_docx_abnt._add_runs_with_markdown(
        p, "See [the docs](https://example.org/d) for details."
    )
    runs = p.runs
    assert "the docs" in "".join(r.text for r in runs)
    # No literal "[the docs]" or "(https" leak.
    assert "[the docs]" not in "".join(r.text for r in runs)
    assert "https://example.org" not in "".join(r.text for r in runs)


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_docx_full_render_does_not_leak_markdown_markers(tmp_path: Path) -> None:
    """End-to-end DOCX render: Markdown markers do not survive in the output."""
    import docx as docx_mod

    import render_docx_abnt

    content = {
        "sections": [
            {
                "id": "05",
                "title": "Síntese",
                "paragraphs": [
                    {
                        "text": "**Convergência forte** sobre *feedback imediato*.",
                        "type": "body",
                    },
                ],
            },
        ],
        "references": [],
    }
    out = tmp_path / "manuscript.docx"
    render_docx_abnt.render_docx_abnt(content, out, title="T")
    doc = docx_mod.Document(str(out))
    body = "\n".join(p.text for p in doc.paragraphs)
    # Markdown markers must NOT appear literally.
    assert "**Convergência forte**" not in body
    assert "*feedback imediato*" not in body
    # The text content survives.
    assert "Convergência forte" in body
    assert "feedback imediato" in body
    # And the bold/italic flags are set on the right runs.
    found_bold = False
    found_italic = False
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.bold and "Convergência forte" in run.text:
                found_bold = True
            if run.italic and "feedback imediato" in run.text:
                found_italic = True
    assert found_bold
    assert found_italic


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_docx_reference_paragraph_does_not_process_markdown(tmp_path: Path) -> None:
    """Reference type in DOCX stays plain — ABNT refs are pre-formatted strings."""
    import docx as docx_mod

    import render_docx_abnt

    content = {
        "sections": [
            {
                "id": "10",
                "title": "Refs",
                "paragraphs": [
                    {"text": "LITERAL **markers** in a ref.", "type": "reference"},
                ],
            },
        ],
        "references": [],
    }
    out = tmp_path / "manuscript.docx"
    render_docx_abnt.render_docx_abnt(content, out, title="T")
    doc = docx_mod.Document(str(out))
    # The literal "**markers**" survives because reference type is plain.
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "**markers**" in body
